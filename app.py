"""Local WeChat article library. Bind only to loopback; no cloud services."""
import base64
from contextlib import contextmanager
import datetime as dt
import email.utils
import hashlib
import html
from html.parser import HTMLParser
import io
import json
import os
from pathlib import Path
import re
import secrets
import sqlite3
import subprocess
import sys
import tempfile
import threading
import time
import unicodedata
import urllib.error
import urllib.parse
import urllib.request
import uuid
import xml.etree.ElementTree as ET
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer

ROOT = Path(__file__).resolve().parent
FROZEN = bool(getattr(sys, 'frozen', False))
DATA = Path(os.environ.get('COLLECTOR_DATA_DIR') or (str(Path(os.environ['LOCALAPPDATA']) / 'Shiwen' / 'data') if FROZEN else str(ROOT / 'data')))
DATA.mkdir(parents=True, exist_ok=True)
DB = DATA / 'library.sqlite3'
PORT = int(os.environ.get('COLLECTOR_PORT', '18765'))
TOKEN = secrets.token_urlsafe(32)
HEADERS = ['公众号','ID','链接','标题','封面','摘要','创建时间','发布时间','阅读','点赞','分享','喜欢','留言','作者','是否原创','文章类型','所属合集','文章内容']
JOBS = {}
LOCK = threading.Lock()

@contextmanager
def db():
    c = sqlite3.connect(DB, timeout=30)
    c.row_factory = sqlite3.Row
    try:
        with c:
            yield c
    finally:
        c.close()

def initialize():
    with db() as c:
        c.executescript('''CREATE TABLE IF NOT EXISTS articles (key TEXT PRIMARY KEY, data TEXT NOT NULL, source TEXT NOT NULL, collected TEXT NOT NULL);
        CREATE TABLE IF NOT EXISTS accounts (name TEXT PRIMARY KEY, feed TEXT NOT NULL DEFAULT '');
        CREATE TABLE IF NOT EXISTS presets (name TEXT PRIMARY KEY, data TEXT NOT NULL);''')

def normalized(s):
    return unicodedata.normalize('NFKC', str(s or '')).casefold().strip()

def canonical_url(url):
    p = urllib.parse.urlsplit(str(url or '').strip())
    if p.scheme not in ('http', 'https') or not p.hostname:
        return ''
    q = urllib.parse.parse_qs(p.query)
    if p.hostname == 'mp.weixin.qq.com':
        keys = ('__biz', 'mid', 'idx')
        if all(q.get(k) for k in keys):
            return 'https://mp.weixin.qq.com/s?' + urllib.parse.urlencode([(k, q[k][0]) for k in keys])
        if p.path.startswith('/s/'):
            return 'https://mp.weixin.qq.com' + p.path.rstrip('/')
    return urllib.parse.urlunsplit((p.scheme, p.netloc, p.path, p.query, ''))

def normalize_source(url):
    url=html.unescape(str(url or '').strip())
    parts=urllib.parse.urlsplit(url)
    if parts.hostname=='mp.weixin.qq.com' and parts.path=='/mp/appmsgalbum':
        params=urllib.parse.parse_qs(parts.query)
        if not params.get('__biz') or not params.get('album_id'):
            raise ValueError('合集链接不完整，请从合集页面复制完整分享链接。')
        return 'https://mp.weixin.qq.com/mp/appmsgalbum?'+urllib.parse.urlencode({'action':'getalbum','__biz':params['__biz'][0],'album_id':params['album_id'][0]})
    return url

def date_text(value):
    if value is None or value == '':
        return ''
    if isinstance(value, (dt.datetime, dt.date)):
        return value.strftime('%Y-%m-%d %H:%M:%S')
    if isinstance(value, (int, float)):
        if value > 1000000000:
            return dt.datetime.fromtimestamp(value, dt.timezone(dt.timedelta(hours=8))).strftime('%Y-%m-%d %H:%M:%S')
        if 20000 < value < 100000:
            return (dt.datetime(1899,12,30) + dt.timedelta(days=value)).strftime('%Y-%m-%d %H:%M:%S')
    s = str(value).strip()
    try:
        d = dt.datetime.fromisoformat(s.replace('Z', '+00:00'))
    except ValueError:
        try:
            d = email.utils.parsedate_to_datetime(s)
        except (ValueError, TypeError):
            m = re.match(r'(\d{4})[年/.-](\d{1,2})[月/.-](\d{1,2})日?(.*)', s)
            if not m:
                return s
            try:
                d = dt.datetime(int(m[1]),int(m[2]),int(m[3]))
                tail = m[4].strip()
                if tail:
                    d = dt.datetime.fromisoformat(d.strftime('%Y-%m-%d') + ' ' + tail)
            except ValueError:
                return s
    if d.tzinfo:
        d = d.astimezone(dt.timezone(dt.timedelta(hours=8))).replace(tzinfo=None)
    return d.strftime('%Y-%m-%d %H:%M:%S')

def ingest(rows, source):
    added = updated = skipped = 0
    with db() as c:
        for incoming in rows:
            row = {h: incoming.get(h, '') for h in HEADERS}
            row = {k: ('' if v is None else v) for k,v in row.items()}
            for key in ('公众号','标题','链接','ID'):
                row[key] = str(row[key]).strip()
            if not row['标题'] or not row['公众号']:
                skipped += 1
                continue
            for field in ('创建时间','发布时间'):
                row[field] = date_text(row[field])
            url = canonical_url(row['链接'])
            identity = url or (row['公众号'] + '\0' + (row['ID'] or row['标题'] + '\0' + row['发布时间']))
            key = hashlib.sha256(identity.encode()).hexdigest()
            previous = c.execute('SELECT data FROM articles WHERE key=?',(key,)).fetchone()
            if not previous and re.fullmatch(r'\d+_\d+',row['ID']):
                equivalent=c.execute("SELECT key,data FROM articles WHERE json_extract(data, '$.公众号')=? AND json_extract(data, '$.ID')=?",(row['公众号'],row['ID'])).fetchone()
                if equivalent:
                    key=equivalent['key']
                    previous=equivalent
            if previous:
                old = json.loads(previous['data'])
                merged = {h: (row[h] if row[h] != '' else old.get(h,'')) for h in HEADERS}
                if merged == old:
                    skipped += 1
                    continue
                row = merged
                updated += 1
            else:
                added += 1
            c.execute('INSERT OR REPLACE INTO articles VALUES (?,?,?,?)', (key,json.dumps(row,ensure_ascii=False),source,dt.datetime.now().isoformat(timespec='seconds')))
            c.execute('INSERT OR IGNORE INTO accounts(name) VALUES (?)',(row['公众号'],))
    return {'added':added,'updated':updated,'skipped':skipped}

def read_xlsx(raw):
    import openpyxl
    w = openpyxl.load_workbook(io.BytesIO(raw),read_only=True,data_only=True)
    rows = []
    for s in w:
        values = iter(s.values)
        columns = [str(x or '').strip() for x in next(values, [])]
        if not all(x in columns for x in ('公众号','标题')):
            continue
        for row in values:
            item = dict(zip(columns,row))
            if any(v not in (None,'') for v in row):
                rows.append(item)
            if len(rows) > 30000:
                raise ValueError('一次最多导入 30000 篇文章，请拆分文件。')
    w.close()
    if not rows:
        raise ValueError('未找到文章。表格第一行必须包含“公众号”和“标题”列。')
    return rows

def keywords(value):
    return [normalized(x) for x in re.split(r'[,，;；\n|]+',str(value or '')) if normalized(x)]

def validate_filter(f):
    for k in ('start','end'):
        if f.get(k):
            dt.date.fromisoformat(f[k])
    if f.get('start') and f.get('end') and f['start'] > f['end']:
        raise ValueError('开始日期不能晚于结束日期。')
    return f

def matches(row, f):
    if f.get('accounts') and row['公众号'] not in f['accounts']:
        return False
    d = row.get('发布时间','')[:10]
    if f.get('start') or f.get('end'):
        try:
            dt.date.fromisoformat(d)
        except ValueError:
            return False
        if f.get('start') and d < f['start'] or f.get('end') and d > f['end']:
            return False
    text = normalized(row.get('标题',''))
    if f.get('scope') == 'content':
        text += '\n' + normalized(row.get('摘要','')) + '\n' + normalized(row.get('文章内容',''))
    inc, exc = keywords(f.get('include')), keywords(f.get('exclude'))
    if inc and not (all(k in text for k in inc) if f.get('mode') == 'all' else any(k in text for k in inc)):
        return False
    if any(k in text for k in exc):
        return False
    if f.get('original') and str(row.get('是否原创')) not in ('原创','是','1','True'):
        return False
    return True

def query(f):
    validate_filter(f)
    with db() as c:
        records = c.execute('SELECT * FROM articles').fetchall()
    result = []
    for rec in records:
        item = json.loads(rec['data'])
        if matches(item,f):
            result.append(dict(item, _key=rec['key'], _source=rec['source'], _collected=rec['collected']))
    return sorted(result,key=lambda x:(x['发布时间'],x['标题']),reverse=True)

class PageParser(HTMLParser):
    def __init__(self):
        super().__init__(convert_charrefs=True)
        self.meta = {}
        self.stack = []
        self.fields = {}
        self.title = []
        self.ignored = 0
    def handle_starttag(self,tag,attrs):
        a = dict(attrs)
        if tag == 'meta':
            self.meta[a.get('property',a.get('name',''))] = a.get('content','')
        if tag not in ('meta','link','img','input','br','hr','source','wbr','area','base','embed','param','track','col'):
            self.stack.append((tag,a.get('id','')))
        if tag in ('p','br','section'):
            for _, key in self.stack:
                if key == 'js_content': self.fields.setdefault(key,[]).append('\n')
    def handle_endtag(self,tag):
        for i in range(len(self.stack)-1,-1,-1):
            if self.stack[i][0] == tag:
                self.stack = self.stack[:i]
                break
    def handle_data(self,data):
        if any(t in ('script','style') for t,_ in self.stack):
            return
        for tag,key in self.stack:
            if key in ('js_content','activity-name','js_name','publish_time','js_author_name'):
                self.fields.setdefault(key,[]).append(data)
            if tag == 'title':
                self.title.append(data)
    def get(self,key):
        return ''.join(self.fields.get(key,[])).strip()

def parse_article(text,url='',fallback=''):
    p = PageParser()
    p.feed(text)
    title = p.get('activity-name') or p.meta.get('og:title','')
    if not title or 'js_content' not in p.fields:
        if any(x in text for x in ('环境异常','访问过于频繁','wappoc_appmsgcaptcha','完成验证')):
            raise ValueError('微信要求人工验证。请在浏览器正常打开文章并完成验证；本工具不会绕过验证。也可导入已保存的文章网页。')
        raise ValueError('未取得文章正文，可能已删除、需要登录或页面结构已变化。')
    def jsvar(name):
        m = re.search(r'\b(?:var|let|const)\s+'+re.escape(name)+r'\s*=\s*[\x22\x27]([^\x22\x27]*)',text)
        return html.unescape(m[1]) if m else ''
    stamp = jsvar('ct') or jsvar('create_time')
    if not stamp:
        m = re.search(r'\bct\s*=\s*(\d{10})', text)
        stamp = m[1] if m else ''
    published = date_text(int(stamp)) if stamp.isdigit() else date_text(p.get('publish_time'))
    account = p.get('js_name') or jsvar('nickname') or fallback
    if not account:
        raise ValueError('未识别到公众号，请指定公众号名称后重试。')
    query_parts = urllib.parse.parse_qs(urllib.parse.urlsplit(url).query)
    mid = jsvar('appmsgid') or jsvar('mid') or query_parts.get('mid',[''])[0]
    idx = jsvar('idx') or query_parts.get('idx',[''])[0]
    article_id = mid+'_'+idx if mid.isdigit() and idx.isdigit() else ''
    return {'公众号':account,'ID':article_id,'链接':url or p.meta.get('og:url',''),'标题':title,'封面':p.meta.get('og:image',''),'摘要':p.meta.get('og:description',''),'发布时间':published,'作者':p.get('js_author_name') or p.meta.get('author',''),'文章内容':p.get('js_content'),'文章类型':'普通图文'}

class SafeRedirect(urllib.request.HTTPRedirectHandler):
    def redirect_request(self,req,fp,code,msg,headers,newurl):
        if urllib.parse.urlsplit(newurl).hostname != 'mp.weixin.qq.com':
            raise ValueError('文章跳转到了非微信域名，已停止读取。')
        return super().redirect_request(req,fp,code,msg,headers,newurl)

def fetch(url, article=False):
    p = urllib.parse.urlsplit(url)
    if p.scheme not in ('http','https') or not p.hostname or p.username or p.password:
        raise ValueError('请输入有效的 http/https 地址。')
    if article and (p.hostname != 'mp.weixin.qq.com' or not p.path.startswith('/s')):
        raise ValueError('只接受 mp.weixin.qq.com/s 开头的微信文章链接。')
    opener = urllib.request.build_opener(SafeRedirect()) if article else urllib.request.build_opener()
    req = urllib.request.Request(url,headers={'User-Agent':'Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 Chrome/131.0.0.0 Safari/537.36'})
    with opener.open(req,timeout=20) as r:
        data = r.read(10*1024*1024+1)
        if len(data)>10*1024*1024:
            raise ValueError('网页超过 10 MB，已停止读取。')
        if article and 'captcha' in r.url:
            raise ValueError('微信要求人工验证，未收集到文章。请在浏览器打开原文完成验证，或导入已保存的文章网页。')
        return data.decode(r.headers.get_content_charset() or 'utf-8',errors='replace')

def parse_feed(text,account):
    if '<!DOCTYPE' in text.upper() or '<!ENTITY' in text.upper():
        raise ValueError('订阅源含不支持的 XML 声明。')
    root = ET.fromstring(text)
    def name(tag):return tag.split('}')[-1]
    rows = []
    for entry in root.iter():
        if name(entry.tag) not in ('item','entry'):
            continue
        fields = {}
        for x in entry:
            k = name(x.tag)
            if k == 'link':
                if x.get('rel','alternate') == 'alternate':
                    fields[k] = x.get('href') or x.text or ''
            else:
                fields[k] = ''.join(x.itertext())
        link = fields.get('link','')
        if urllib.parse.urlsplit(link).hostname != 'mp.weixin.qq.com':
            continue
        def plain(s):
            return html.unescape(re.sub('<[^>]*>',' ',s)).strip()
        rows.append({'公众号':account,'标题':plain(fields.get('title','')),'链接':link,'发布时间':date_text(fields.get('pubDate') or fields.get('published') or fields.get('updated','')),'作者':plain(fields.get('creator') or fields.get('author','')),'摘要':plain(fields.get('description') or fields.get('summary','')),'文章内容':plain(fields.get('encoded') or fields.get('content',''))})
    return rows

def album_batches(url, account):
    """Read public album metadata using its native pagination."""
    parts=urllib.parse.urlsplit(html.unescape(url))
    params=urllib.parse.parse_qs(parts.query)
    if parts.hostname!='mp.weixin.qq.com' or parts.path!='/mp/appmsgalbum' or not params.get('__biz') or not params.get('album_id'):
        raise ValueError('请输入完整的微信公众号合集分享链接。')
    page=fetch(url)
    meta=page.split('window.cgiData =',1)[-1].split('articleList:',1)[0]
    def field(k):
        m=re.search(r'\b'+k+r"\s*:\s*'((?:\\.|[^'\\])*)'",meta)
        return html.unescape(m[1].replace('\\x26','&').replace("\\'","'")) if m else ''
    real_account=field('nick_name')
    title=field('title')
    if not real_account or not title:
        raise ValueError('未能读取合集，微信可能要求验证，或该合集已不可用。')
    if normalized(real_account)!=normalized(account):
        raise ValueError('该合集属于“'+real_account+'”，请将公众号名称填写为这个名称。')
    params={'action':'getalbum','__biz':params['__biz'][0],'album_id':params['album_id'][0],'count':'20','is_reverse':'1','f':'json'}
    visited=set()
    for _ in range(100):
        raw=fetch('https://mp.weixin.qq.com/mp/appmsgalbum?'+urllib.parse.urlencode(params))
        try:data=json.loads(raw)
        except json.JSONDecodeError:raise ValueError('微信未返回合集数据，可能需要验证。已收集的文章已保留。')
        if data.get('base_resp',{}).get('ret') not in (0,'0'):
            raise ValueError('微信合集请求未成功：'+str(data.get('base_resp',{}).get('ret','未知')))
        result=data.get('getalbum_resp',{})
        items=result.get('article_list',[])
        if isinstance(items,dict):items=[items]
        if not items:return
        rows=[]
        for x in items:
            link=html.unescape(x.get('url',''))
            if urllib.parse.urlsplit(link).hostname!='mp.weixin.qq.com':continue
            stamp=x.get('create_time','')
            rows.append({'公众号':real_account,'ID':str(x.get('msgid',''))+'_'+str(x.get('itemidx','')),'链接':link,'标题':html.unescape(x.get('title','')),'封面':x.get('cover_img_1_1',''),'发布时间':date_text(int(stamp)) if str(stamp).isdigit() else '', '所属合集':'#'+title,'文章类型':'普通图文'})
        yield rows
        if str(result.get('continue_flag','0'))!='1':return
        cursor=(str(items[-1].get('msgid','')),str(items[-1].get('itemidx','')))
        if cursor in visited:raise ValueError('合集分页未推进，已停止并保留已收集文章。')
        visited.add(cursor)
        params.update(begin_msgid=cursor[0],begin_itemidx=cursor[1])
        time.sleep(0.6)
    raise ValueError('单次采集达到 100 页上限，已保留已读取文章。')

def collection_worker(job_id,payload):
    tasks = []
    if payload.get('kind') == 'links':
        urls = re.findall(r'https?://[^\s<>\x22]+',payload.get('urls',''))
        urls = list(dict.fromkeys(html.unescape(x).rstrip('，。；') for x in urls))
        tasks = [('link',u,payload.get('account','')) for u in urls[:100]]
    else:
        with db() as c:
            accounts = [dict(x) for x in c.execute('SELECT * FROM accounts')]
        chosen = payload.get('accounts',[])
        tasks = [('feed',a['feed'],a['name']) for a in accounts if a['feed'] and (not chosen or a['name'] in chosen)]
    with LOCK:
        JOBS[job_id].update(total=len(tasks))
    if not tasks:
        with LOCK:
            JOBS[job_id].update(done=True,errors=['没有可采集的来源。请粘贴微信文章链接，或为公众号配置合集 / RSS 地址。'])
        return
    for i,(kind,url,account) in enumerate(tasks):
        try:
            is_album=kind!='link' and urllib.parse.urlsplit(url).hostname=='mp.weixin.qq.com' and urllib.parse.urlsplit(url).path=='/mp/appmsgalbum'
            if is_album:
                batches=album_batches(url,account)
                source='微信合集：'+account
            else:
                text = fetch(url,article=kind=='link')
                rows = [parse_article(text,url,account)] if kind=='link' else parse_feed(text,account)
                if not rows:
                    raise ValueError('订阅源没有可识别的微信文章。请确认地址为 RSS/Atom，且文章链接指向微信。')
                batches=[rows]
                source='微信链接' if kind=='link' else '订阅源：'+account
            for rows in batches:
                result=ingest(rows,source)
                with LOCK:
                    for k in result:JOBS[job_id][k]+=result[k]
        except Exception as e:
            with LOCK:
                JOBS[job_id]['errors'].append((account or url[:90])+'：'+str(e))
        with LOCK:JOBS[job_id]['completed']=i+1
        if i < len(tasks)-1:time.sleep(2)
    with LOCK:JOBS[job_id]['done']=True

def export_xlsx(rows):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.worksheet.table import Table, TableStyleInfo
    from openpyxl.utils import get_column_letter
    wb=Workbook(); sheet=wb.active; sheet.title='Sheet1'
    sheet.append(HEADERS)
    for row in rows:
        values=[]
        for h in HEADERS:
            value=row.get(h, '')
            if isinstance(value,str):
                value=re.sub(r'[\x00-\x08\x0b\x0c\x0e-\x1f]', '', value)
                value=value.encode('utf-16-le',errors='replace')[:65534].decode('utf-16-le',errors='ignore')
                if h in ('创建时间','发布时间'):
                    try:value=dt.datetime.strptime(value,'%Y-%m-%d %H:%M:%S')
                    except ValueError:pass
            values.append(value)
        sheet.append(values)
        for cell in sheet[sheet.max_row]:
            # Treat article text as literal text, including titles starting with '='.
            if isinstance(cell.value,str):cell.data_type='s'
            cell.font=Font(name='Microsoft YaHei',size=11)
            cell.alignment=Alignment(vertical='center',wrap_text=True)
            if isinstance(cell.value,dt.datetime):cell.number_format='yyyy-mm-dd hh:mm:ss'
        sheet.row_dimensions[sheet.max_row].height=60
    for cell in sheet[1]:
        cell.font=Font(name='Microsoft YaHei',bold=True,color='FFFFFF',size=11)
        cell.fill=PatternFill('solid',fgColor='183B36')
    sheet.row_dimensions[1].height=30
    widths=[28,22,40,70,40,45,23,23,12,12,12,12,12,18,12,18,28,60]
    for col,width in enumerate(widths,1):sheet.column_dimensions[get_column_letter(col)].width=width
    sheet.freeze_panes='A2'
    sheet.auto_filter.ref=f'A1:R{sheet.max_row}'
    if rows:
        table=Table(displayName='Articles',ref=f'A1:R{sheet.max_row}')
        table.tableStyleInfo=TableStyleInfo(name='TableStyleMedium2',showRowStripes=True)
        sheet.add_table(table)
    output=io.BytesIO();wb.save(output);wb.close()
    return output.getvalue()

class Handler(BaseHTTPRequestHandler):
    def log_message(self,*args):pass
    def respond(self,obj,status=200):
        raw=json.dumps(obj,ensure_ascii=False,default=str).encode()
        self.send_response(status)
        self.send_header('Content-Type','application/json; charset=utf-8')
        self.send_header('Content-Length',str(len(raw)))
        self.send_header('Cache-Control','no-store')
        self.end_headers()
        self.wfile.write(raw)
    def allowed(self):
        host=self.headers.get('Host','')
        return host in (f'127.0.0.1:{PORT}',f'localhost:{PORT}')
    def do_GET(self):
        if not self.allowed():return self.respond({'error':'仅允许本机访问'},403)
        path=urllib.parse.urlsplit(self.path).path
        if path=='/health':return self.respond({'app':'wechat-collector','version':1,'desktop':FROZEN,'instance':TOKEN[:12]})
        if path in ('/','/app.js','/style.css'):
            name={'/':'index.html','/app.js':'app.js','/style.css':'style.css'}[path]
            raw=(ROOT/'dist'/name).read_bytes()
            if name=='index.html':raw=raw.replace(b'__TOKEN__',TOKEN.encode())
            self.send_response(200)
            self.send_header('Content-Type',{'html':'text/html; charset=utf-8','js':'text/javascript; charset=utf-8','css':'text/css; charset=utf-8'}[name.split('.')[-1]])
            self.send_header('Content-Length',str(len(raw)))
            self.send_header('Cache-Control','no-store')
            self.send_header('X-Frame-Options','DENY')
            self.send_header('Content-Security-Policy',"default-src 'self'; style-src 'self'; script-src 'self'; connect-src 'self'; img-src 'self' data:; frame-ancestors 'none'; base-uri 'none'")
            self.end_headers();self.wfile.write(raw);return
        return self.respond({'error':'未找到'},404)
    def do_POST(self):
        if not self.allowed() or self.headers.get('X-Collector-Token')!=TOKEN:
            return self.respond({'error':'页面授权已过期，请刷新页面'},403)
        if self.headers.get('Origin') not in (None,f'http://127.0.0.1:{PORT}',f'http://localhost:{PORT}'):
            return self.respond({'error':'不允许跨站访问'},403)
        try:
            length=int(self.headers.get('Content-Length','0'))
            if length>24*1024*1024:raise ValueError('文件过大，请控制在 16 MB 内。')
            p=json.loads(self.rfile.read(length) or b'{}')
            route=urllib.parse.urlsplit(self.path).path
            if route=='/api/shutdown' and FROZEN:
                with LOCK:
                    if any(not job['done'] for job in JOBS.values()) and not p.get('force'):
                        raise ValueError('正在收集文章，请等待采集完成后退出。')
                self.respond({'ok':True})
                threading.Thread(target=self.server.shutdown,daemon=True).start()
                return
            if route=='/api/state':
                with db() as c:
                    accounts=[dict(x) for x in c.execute('SELECT * FROM accounts ORDER BY name')]
                    counts={x[0]:x[1] for x in c.execute("SELECT json_extract(data, '$.公众号'), count(*) FROM articles GROUP BY 1")}
                    for account in accounts:
                        account['articleCount']=counts.get(account['name'],0)
                        account['feed']=normalize_source(account['feed'])
                    presets=[{'name':x['name'],'filter':json.loads(x['data'])} for x in c.execute('SELECT * FROM presets ORDER BY name')]
                    total=c.execute('SELECT count(*) FROM articles').fetchone()[0]
                return self.respond({'accounts':accounts,'presets':presets,'total':total})
            if route=='/api/query':
                rows=query(p.get('filter',{}));page=max(1,int(p.get('page',1)));size=40
                return self.respond({'rows':rows[(page-1)*size:page*size],'total':len(rows),'page':page,'pages':max(1,(len(rows)+size-1)//size)})
            if route=='/api/import':
                raw=base64.b64decode(p['data'],validate=True)
                rows=read_xlsx(raw) if p.get('name','').lower().endswith('.xlsx') else [parse_article(raw.decode('utf-8-sig'),fallback=p.get('account',''))]
                return self.respond(ingest(rows,'导入：'+Path(p.get('name','file')).name))
            if route=='/api/preset':
                name=str(p.get('name','')).strip()[:80]
                if not name:raise ValueError('请输入方案名称。')
                f=validate_filter(p.get('filter',{}))
                with db() as c:
                    c.execute('INSERT OR REPLACE INTO presets VALUES (?,?)',(name,json.dumps(f,ensure_ascii=False)))
                return self.respond({'ok':True})
            if route=='/api/account':
                name=str(p.get('name','')).strip()[:120]
                if not name:raise ValueError('请输入公众号名称。')
                feed=normalize_source(p.get('feed',''))
                if feed and urllib.parse.urlsplit(feed).scheme not in ('http','https'):raise ValueError('订阅地址必须以 http:// 或 https:// 开头。')
                with db() as c:c.execute('INSERT OR REPLACE INTO accounts VALUES (?,?)',(name,feed))
                return self.respond({'ok':True})
            if route=='/api/collect':
                if p.get('kind') == 'links':
                    links = re.findall(r'https?://[^\s<>\x22]+',p.get('urls',''))
                    if not links:
                        raise ValueError('请粘贴有效的微信文章链接。')
                    if len(set(links)) > 100:
                        raise ValueError('每次最多采集 100 条链接，请分批收集。')
                with LOCK:
                    if any(not j['done'] for j in JOBS.values()):raise ValueError('已有采集任务进行中，请等待完成。')
                    jid=uuid.uuid4().hex
                    JOBS[jid]={'done':False,'total':0,'completed':0,'added':0,'updated':0,'skipped':0,'errors':[]}
                threading.Thread(target=collection_worker,args=(jid,p),daemon=True).start()
                return self.respond({'id':jid})
            if route=='/api/job':
                with LOCK:r=dict(JOBS.get(p.get('id'),{'done':True,'errors':['任务不存在，请重新采集。']}))
                return self.respond(r)
            if route=='/api/export':
                rows=query(p.get('filter',{}))
                if p.get('keys'):
                    selected=set(p['keys']);rows=[r for r in rows if r['_key'] in selected]
                raw=export_xlsx(rows)
                self.send_response(200)
                self.send_header('Content-Type','application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
                self.send_header('Content-Disposition','attachment; filename="wechat-articles.xlsx"')
                self.send_header('Content-Length',str(len(raw)))
                self.end_headers();self.wfile.write(raw);return
            return self.respond({'error':'接口不存在'},404)
        except (ValueError,KeyError,TypeError,ET.ParseError) as e:
            return self.respond({'error':str(e)},400)
        except Exception as e:
            return self.respond({'error':'操作未完成：'+str(e)},500)

if __name__=='__main__':
    initialize()
    if len(sys.argv)>2 and sys.argv[1]=='--import':
        print(json.dumps(ingest(read_xlsx(Path(sys.argv[2]).read_bytes()),'导入：'+Path(sys.argv[2]).name),ensure_ascii=False))
    else:
        print(f'http://127.0.0.1:{PORT}',flush=True)
        ThreadingHTTPServer(('127.0.0.1',PORT),Handler).serve_forever()
